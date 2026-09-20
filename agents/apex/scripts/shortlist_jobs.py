#!/usr/bin/env python3
"""Deterministic, offline profile matching and XLSX export. No LLM or network calls."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timedelta, timezone
import hashlib
from html import unescape
import json
import math
from pathlib import Path
import re
import sqlite3
import sys
import tempfile
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

ROOT = Path(__file__).resolve().parents[3]
UTC = timezone.utc
# Title signals are intentionally narrower than general skills in a job description.
# Each family needs both a vacancy-title signal and an applicant evidence signal.
FAMILIES = {
    'Social protection and cash': (r'social (?:protection|policy|affairs)|cash.based|cash transfer|emergency cash', r'social protection|cash transfer|\bHOPE\b'),
    'Monitoring and evaluation': (r'monitoring|evaluat(?:ion|or)|\bM&E\b|\bMEAL\b', r'monitoring|evaluation|post.distribution'),
    'Data and information management': (r'data (?:management|analyst|analysis|governance|and monitoring)|information management|knowledge management', r'\bHOPE\b|data (?:management|cleaning|quality)|Power BI|Python'),
    'Humanitarian response': (r'humanitarian|emergency (?:response|programme|program)|displacement|resilience', r'humanitarian|emergency|displacement'),
    'Grants and partnerships': (r'\bgrants?\b|partnership|resource mobili[sz]|external relations|assurance', r'grant|partnership|donor|HACT'),
    'Programme management': (r'program(?:me)? (?:management|manager|officer|specialist|coordinat|analyst)|project (?:manager|management|officer|specialist|coordinat)|work programme', r'programme|program management|project management|coordination'),
    'Administration and operations': (r'administrative officer|assistant to team|operations (?:officer|manager|specialist)', r'administrati|accounting|logistics|operations'),
    'Training and capacity development': (r'training|capacity (?:building|development)|learning specialist', r'training|capacity.building|capacity.development'),
}
SPECIALIST_EXCLUSIONS = re.compile(r'\b(?:internship|intern|trainee|student|surgeon|nurse|physicist|engineer|lawyer|translator|interpreter|data scientist|cybersecurity|medical officer|air operations|procurement officer|aviation|aircraft|air navigation|CNS/ATM|fire and rescue)\b', re.I)
LANGUAGES = ('English', 'Japanese', 'French', 'Spanish', 'Arabic', 'Portuguese', 'Russian', 'Chinese', 'German', 'Greek', 'Korean')


def plain(value):
    return re.sub(r'\s+', ' ', unescape(re.sub(r'<[^>]+>', ' ', str(value or '')))).strip()


def section(text, name):
    """Read a canonical h2 section while ignoring headings in fenced blocks."""
    found = []; active = False; fence = None
    for line in text.splitlines():
        marker = re.match(r'^ {0,3}(`{3,}|~{3,})(.*)$', line)
        if fence:
            if marker and marker[1][0] == fence[0] and len(marker[1]) >= fence[1] and not marker[2].strip():
                fence = None
            if active: found.append(line)
            continue
        if marker:
            fence = (marker[1][0], len(marker[1]))
        if line.startswith('## '):
            if active: break
            active = line.strip() == '## ' + name
            continue
        if active: found.append(line)
    return '\n'.join(found).strip()


def read_profile(path, feedback=None, overrides=None):
    text = Path(path).read_text(encoding='utf-8')
    history = section(text, 'USER_JOB_HISTORY_TEXT')
    if not history:
        raise ValueError('Profile needs a populated USER_JOB_HISTORY_TEXT section')
    approved = section(Path(feedback).read_text(encoding='utf-8'), 'APPROVED_UPDATES') if feedback else ''
    # Never mine the target JD, generated admin profile, or unresolved feedback for skills.
    evidence = history + '\n' + approved
    lines = [plain(x) for x in evidence.splitlines() if plain(x)]
    families = {}
    for name, (_, signal) in FAMILIES.items():
        snippets = [x for x in lines if re.search(signal, x, re.I)]
        if snippets: families[name] = snippets[0][:300]
    languages = {}
    language_lines = [x for x in history.splitlines() if re.match(r'\s*Languages?:', x, re.I)]
    for line in language_lines:
        for language in LANGUAGES:
            m = re.search(r'\b' + language + r'\s*\(([^)]+)\)', line, re.I)
            if m: languages[language.lower()] = m[1].lower()
    settings = {'languages': languages, 'citizenships': [], 'national_countries': ['JP'],
                'work_authorized_countries': [], 'disabled_families': []}
    if overrides:
        update = json.loads(Path(overrides).read_text(encoding='utf-8'))
        unknown = set(update) - set(settings)
        if unknown: raise ValueError('Unknown profile override keys: ' + ', '.join(sorted(unknown)))
        settings.update(update)
    for key in ('citizenships', 'national_countries', 'work_authorized_countries', 'disabled_families'):
        if not isinstance(settings[key], list): raise ValueError(key + ' must be a list')
    if not isinstance(settings['languages'], dict): raise ValueError('languages must be a mapping')
    settings['languages'] = {k.lower(): str(v).lower() for k, v in settings['languages'].items()}
    families = {k:v for k,v in families.items() if k not in settings['disabled_families']}
    return {**settings, 'families': families, 'evidence_sha256': hashlib.sha256(evidence.encode()).hexdigest()}


def snapshot(database):
    """Read one consistent transaction and only the columns needed for screening."""
    database = Path(database).resolve(strict=True)
    with sqlite3.connect(database.as_uri() + '?mode=ro', uri=True) as db:
        db.row_factory = sqlite3.Row
        db.execute('PRAGMA query_only=ON'); db.execute('BEGIN')
        tables = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if 'jobs' not in tables: raise ValueError('Database has no jobs table')
        cols = {r[1] for r in db.execute('PRAGMA table_info(jobs)')}
        required = {'job_key','title','status','apply_url','description'}
        if required - cols: raise ValueError('Missing jobs columns: '+', '.join(sorted(required-cols)))
        names = ['job_key','source_id','org_id','external_id','title','location','employment_type','status',
                 'closes_at','closes_at_local','closes_tz','apply_url','source_url','description','raw_json',
                 'last_seen_at','source_latest_observed_at','source_listed_current','trusted_current',
                 'duplicate_of_job_key','stale_current','detail_quality_status']
        select = ['j.'+n for n in names if n in cols]
        join = ''
        if 'vacancy_classifications' in tables:
            ccols = {r[1] for r in db.execute('PRAGMA table_info(vacancy_classifications)')}
            for n in ['grade_code','grade_family','standard_scope','national_international','country_iso2','country','contract_category','work_modality']:
                if n in ccols: select.append('c.'+n)
            join = ' LEFT JOIN vacancy_classifications c ON c.vacancy_id=j.job_key'
        rows = [dict(r) for r in db.execute('SELECT '+','.join(select)+' FROM jobs j'+join+" WHERE lower(j.status)='open'")]
    return rows


def parse_date(value):
    if not value: return None
    s = str(value).strip()
    try: return datetime.fromisoformat(s.replace('Z','+00:00'))
    except ValueError: pass
    for fmt in ('%d-%b-%Y','%d/%b/%Y, %I:%M:%S %p','%d/%b/%Y, %H:%M:%S','%B %d, %Y','%d %B %Y','%d/%m/%Y'):
        try: return datetime.strptime(s,fmt)
        except ValueError: pass
    return None


def closing(row, now):
    """Return known future cutoff, or a future date with precision explicitly marked."""
    value = row.get('closes_at') or row.get('closes_at_local')
    dt = parse_date(value)
    if dt is None: return None, 'missing_or_invalid_deadline'
    timed = bool(re.search(r'(?:T|\s)\d{1,2}:\d{2}', str(value)))
    if timed and dt.tzinfo is None and row.get('closes_tz'):
        try: dt = dt.replace(tzinfo=ZoneInfo(row['closes_tz']))
        except (ZoneInfoNotFoundError, ValueError): pass
    if timed and dt.tzinfo is not None:
        dt = dt.astimezone(UTC)
        if dt <= now: return None, 'past_deadline'
        return {'sort':dt,'date':dt.replace(tzinfo=None),'precision':'UTC instant','published':str(value)+((' '+row['closes_tz']) if row.get('closes_tz') else '')}, None
    # A date with unknown cutoff is never kept on its closing day or afterward.
    if dt.date() <= now.date(): return None, 'past_or_uncertain_same_day_deadline'
    if timed:
        # Unknown local clock: require even UTC+14 interpretation to be in future.
        if dt.replace(tzinfo=timezone(timedelta(hours=14))).astimezone(UTC) <= now:
            return None, 'uncertain_local_cutoff'
    return {'sort':datetime.combine(dt.date(),datetime.min.time(),UTC),
            'date':datetime.combine(dt.date(),datetime.min.time()),'precision':'Date only; exact cutoff unverified',
            'published':str(value)+((' '+row['closes_tz']) if row.get('closes_tz') else '')}, None


def recruitment(row, profile):
    title=plain(row.get('title')); location=plain(row.get('location')); body=plain(row.get('description'))
    grade=str(row.get('grade_code') or '')
    japan=row.get('country_iso2')=='JP' or bool(re.search(r'\b(Japan|Tokyo|Osaka|Kobe|Yokohama)\b',location,re.I))
    remote=bool(re.search(r'\bremote\b|home[ -]based|work from home',title+' '+location,re.I))
    if re.search(r'not remote|remote\s*[:=]?\s*no',title+' '+location,re.I): remote=False
    local=bool(re.search(r'^national\b|\bnational (?:project officer|consultant|consultancy|officer|specialist|expert|coordinator)\b|\bnationals\s*$|nationals[ -]only|\bnationally recruited\b|\blocally recruited\b|\bNPSA[- ]?\d|\bLICA[- ]?\d',title,re.I)
               or re.match(r'^(?:NO[ -]?[A-D1-4]|G[ -]?\d|NPSA|LICA)',grade,re.I)
               or row.get('standard_scope') in ('National / Local','Local','Local / National')
               or row.get('national_international') in ('national','local','unv_national'))
    explicit_local=re.search(r'(?:subject to local recruitment|open only to nationals|only (?:open|available) to nationals|must be (?:a )?(?:national|citizen) of (?!a country other)|national consultant contract|Recruitment Type:\s*Local Recruitment)',body,re.I)
    local=local or bool(explicit_local)
    issues=[]
    if local and not (japan and 'JP' in profile['national_countries']): return None, remote, ['National/local recruitment outside Japan']
    if japan:
        if local and 'JP' not in profile['citizenships']: issues.append('Japanese citizenship/local eligibility not established in profile')
        if re.search(r'commuting area|resident|residen[ct]e',body,re.I):issues.append('Confirm Japan residence/local commuting-area requirement')
        return 'Japan national/local exception' if local else 'Japan-based role',remote,issues
    if remote:
        if re.search(r'right to work|work authori[sz]ation|must (?:reside|be based)|eligible to work',body,re.I):issues.append('Remote role has location/work-authorization conditions; inspect notice')
        return 'Remote',True,issues
    if re.match(r'^(?:P[ -]?[1-5]|D[ -]?[12]|IICA|IPSA)',grade,re.I) or row.get('standard_scope') in ('International','International / Field','External international volunteer pathway'):
        scope='International'
    elif re.search(r'international (?:consultant|recruitment)|qualified international candidates|nationals of a country other than',title+' '+body,re.I):scope='International'
    else:
        scope='Recruitment scope unconfirmed';issues.append('International recruitment cannot be established from stored fields')
    if re.search(r'member countr|participating states|citizenship|korean nationals|nationality.*restrict',body,re.I):
        issues.append('Check restricted-nationality/member-country eligibility')
    return scope,False,issues


def language_checks(text, profile):
    """Conservative explicit-language rule. Ambiguous alternatives become review flags."""
    gaps=[]; review=[]
    # Required clauses only: avoid generic working-language boilerplate and desirable lists.
    chunks=re.split(r'(?<=[.;])\s+|\n',text)
    for chunk in chunks:
        if not re.search(r'\brequired\b|\bessential\b|\bmandatory\b',chunk,re.I):continue
        if not re.search(r'fluenc|fluent|proficien|knowledge|language',chunk,re.I):continue
        names=[lang for lang in LANGUAGES if re.search(r'\b'+lang+r'\b',chunk,re.I)]
        if not names:continue
        if len(chunk)>450 or re.search(r'desirable|an asset|advantage|preferred',chunk,re.I):
            review.append('Language requirement needs manual reading: '+chunk[:260]);continue
        def level(value):
            if re.search(r'fluent|native|professional|advanced|proficien|expert',value,re.I):return 4
            if re.search(r'intermediate|working',value,re.I):return 3
            if re.search(r'basic|limited|beginner',value,re.I):return 2
            return 0
        minimum=2 if re.search(r'limited|basic',chunk,re.I) else 3 if re.search(r'intermediate|working knowledge',chunk,re.I) else 4
        unsupported=[lang for lang in names if level(profile['languages'].get(lang.lower(),''))<minimum]
        if unsupported:
            if re.search(r'\bor\b|one of|either',chunk,re.I):
                if len(unsupported)==len(names):gaps.append('Required language alternative unsupported: '+', '.join(names))
                else: review.append('Confirm language alternatives: '+chunk[:240])
            else:gaps.append('Required language not evidenced at requested level: '+', '.join(unsupported))
    return list(dict.fromkeys(gaps)),list(dict.fromkeys(review))


def evaluate(row, profile, now):
    if row.get('status','').lower()!='open':return None,'not_open'
    if row.get('duplicate_of_job_key'):return None,'duplicate'
    if row.get('stale_current'):return None,'stale'
    if 'source_listed_current' in row and not row['source_listed_current']:return None,'not_currently_listed'
    deadline,reason=closing(row,now)
    if reason:return None,reason
    title=plain(row.get('title'));body=plain(row.get('description'))
    if SPECIALIST_EXCLUSIONS.search(title):return None,'specialist_or_career_stage_mismatch'
    if re.search(r'online volunteering|unpaid|pro.bono',title+' '+body[:500],re.I):return None,'unpaid'
    families=[name for name,(signal,_) in FAMILIES.items() if name in profile['families'] and re.search(signal,title,re.I)]
    if not families:return None,'no_profile_title_overlap'
    scope,remote,issues=recruitment(row,profile)
    if scope is None:return None,'national_local_outside_japan'
    gaps,lang_review=language_checks(body,profile)
    if gaps:return None,'mandatory_language_gap'
    issues.extend(lang_review)
    if len(body)<500:issues.append('Incomplete stored job description')
    if re.search(r'only.*internal|internal candidates only|open to tier[s]? 0,? 1.*2|open to internal applicants only',title+' '+body,re.I):
        issues.append('Internal-candidate restriction may apply; current staff eligibility unverified')
    if re.search(r'\bP[ -]?5\b|\bD[ -]?[12]\b|\bchief\b|\bdirector\b',title+' '+str(row.get('grade_code')),re.I):issues.append('Senior leadership scope and required specialist years need review')
    # Preserve concrete mandatory criteria for human assessment, not just broad title overlap.
    criteria=[s.strip() for s in re.split(r'(?<=[.;])\s+|\n',body) if re.search(r'\brequired\b|\bat least\b|\bminimum\b|\bmust\b',s,re.I)
              and not re.search(r'vaccin|medical clearance|bank account|sexual|background check|inoculat',s,re.I)]
    criteria_text='\n'.join(s[:450] for s in criteria[:6])[:1400]
    if not criteria_text:issues.append('No reliable mandatory-criteria extract; read full notice')
    raw={}
    try:raw=json.loads(row.get('raw_json') or '{}')
    except (ValueError,TypeError):pass
    if not isinstance(raw,dict):raw={}
    url=raw.get('_pageup_detail_url') or raw.get('_taleo_detail_url') or row.get('source_url') or row.get('apply_url')
    if row.get('source_id')=='un_inspira' and row.get('external_id'):url='https://careers.un.org/jobSearchDescription/'+str(row['external_id'])+'?language=en'
    grade=str(row.get('grade_code') or 'Not published')
    if grade.startswith('EXP'):grade='Not published'
    if re.search(r'consultant|consultancy',title,re.I) or grade in ('CON','UG'):grade='Consultant (ungraded)'
    why='; '.join(families)+'. Matching profile evidence: '+' / '.join(profile['families'][f] for f in families[:2])
    roster=bool(re.search(r'\broster\b|talent pool',title,re.I))
    return {'key':row['job_key'],'closing':deadline['date'],'precision':deadline['precision'],
            'published_deadline':deadline['published'],'organization':row.get('org_id') or row.get('source_id',''),
            'title':title,'location':'Remote' if remote else plain(row.get('location')) or 'Not published',
            'grade':grade,'scope':scope,'families':', '.join(families),'why':why,
            'checks':'\n'.join(dict.fromkeys(issues)) or 'Verify specialised experience, education and years against the full notice.',
            'requirements':criteria_text,'apply_url':row.get('apply_url') or url,'source_url':url,
            'last_seen':row.get('last_seen_at') or 'Not recorded','contract':row.get('employment_type') or row.get('contract_category') or 'Not published',
            'group':'Open roster calls' if roster else 'Review eligibility' if issues else 'Potential matches',
            'sort':deadline['sort']},None


def screen(rows, profile, now):
    kept=[];excluded=Counter();seen=set()
    for row in sorted(rows,key=lambda r:r['job_key']):
        record,reason=evaluate(row,profile,now)
        if reason:excluded[reason]+=1;continue
        # Source ID/key is authoritative; never merge unrelated jobs sharing a title or generic apply URL.
        if record['key'] in seen:excluded['duplicate']+=1;continue
        seen.add(record['key']);kept.append(record)
    kept.sort(key=lambda r:(r['sort'],r['title'],r['key']))
    return kept,dict(excluded)


def export_xlsx(records, profile, now, excluded, output):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError('Install Excel support: python -m pip install -r agents/apex/job_shortlist_requirements.txt') from exc
    wb=Workbook();wb.remove(wb.active)
    columns=[('Closing date / UTC time','closing',23),('Organization','organization',22),('Title','title',57),('Duty station','location',27),
             ('Grade / contract level','grade',23),('Recruitment scope','scope',31),('Matching profile areas','families',39),
             ('Why included','why',72),('Checks before applying','checks',72),('Apply URL','apply_url',23),
             ('Mandatory criteria extract','requirements',95),('Official vacancy URL','source_url',24),('Contract','contract',26),
             ('Deadline precision','precision',30),('Published deadline','published_deadline',33),('Database last seen','last_seen',30),('Vacancy key','key',38)]
    for index,name in enumerate(('Potential matches','Review eligibility','Open roster calls')):
        sh=wb.create_sheet(name);sh.sheet_view.showGridLines=False
        sh.append([name]);sh['A1'].font=Font(name='Calibri',size=16,bold=True,color='234760')
        sh.append([f'As of {now:%Y-%m-%d %H:%M UTC}. Rule-based relevance only; all rows are open/currently listed in the database.'])
        sh.append(['Timed deadlines are UTC. Date-only cutoffs are unverified. No live portal check or application submission.'])
        sh.append([c[0] for c in columns]);selected=[r for r in records if r['group']==name]
        for r in selected:
            sh.append([r.get(key,'') for _,key,_ in columns])
            for c,(_,key,_) in zip(sh[sh.max_row],columns):
                # Stored vacancy text is untrusted spreadsheet input: never create formulas.
                if isinstance(c.value,str):c.data_type='s'
                c.alignment=Alignment(vertical='top',wrap_text=True)
                c.font=Font(name='Calibri',size=11)
                if key in ('apply_url','source_url'):
                    if str(c.value).startswith(('https://','http://')):
                        c.hyperlink=c.value;c.value='Apply' if key=='apply_url' else 'Official notice';c.style='Hyperlink'
                if key=='closing':c.number_format='yyyy-mm-dd hh:mm' if r['precision']=='UTC instant' else 'yyyy-mm-dd'
            line_count=max(sum(max(1,math.ceil(len(line)/(width*0.9))) for line in str(r.get(key,'')).split('\n')) for _,key,width in columns)
            sh.row_dimensions[sh.max_row].height=min(400,max(65,line_count*15+12))
        for i,(_,_,width) in enumerate(columns,1):sh.column_dimensions[sh.cell(4,i).column_letter].width=width
        for cell in sh[4]:cell.fill=PatternFill('solid',fgColor='234760');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='center')
        sh.row_dimensions[4].height=32;sh.freeze_panes='D5'
        if selected:
            tab=Table(displayName='Vacancies'+str(index),ref=f'A4:Q{sh.max_row}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);sh.add_table(tab)
        else:sh.append(['No matching records in this category.'])
    sh=wb.create_sheet('Method and coverage');sh.sheet_view.showGridLines=False
    notes=[('Rule engine','Pure Python. No LLM calls, network access or database writes. Title signals must match capabilities extracted from applicant evidence.'),
           ('Profile boundary','Uses only USER_JOB_HISTORY_TEXT and, if supplied, APPROVED_UPDATES. Target JD, generated admin text and unresolved assertions are excluded.'),
           ('Scope','International roles, national/local positions in Japan and remote roles. National-only roles outside Japan are excluded even if remote.'),
           ('Eligibility review','Uncertain recruitment restrictions, language alternatives and other detected gaps are separated for review. Rules cannot establish complete eligibility or semantic equivalence.'),
           ('Open status','Requires status=open, no duplicate/stale flag, current listing when the field exists, and an unexpired deadline. Unknown deadlines and uncertain same-day local/date-only deadlines are excluded.'),
           ('Limits','The script does not infer qualifying years, degree equivalence, specialist expertise, residence, citizenship, or current staff status. Inspect mandatory criteria and the full notice.'),
           ('Date sorting','Each vacancy sheet is independently sorted by closing date, ascending. Date-only values are sorted at the start of their calendar day without claiming midnight is the cutoff.'),
           ('Run time',now.isoformat()),('Profile evidence SHA-256',profile['evidence_sha256']),('Included records',len(records)),
           ('Languages',json.dumps(profile['languages'],ensure_ascii=False)),('Exclusion counts',json.dumps(excluded,sort_keys=True))]
    for name,detail in notes:sh.append([name,detail])
    sh.column_dimensions['A'].width=28;sh.column_dimensions['B'].width=115
    for row in sh:
        sh.row_dimensions[row[0].row].height=60
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=11)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix='.xlsx',dir=output.parent,delete=False) as tmp:temp=Path(tmp.name)
    try:wb.save(temp);temp.replace(output)
    finally:temp.unlink(missing_ok=True)


def main(argv=None):
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--database',type=Path,default=ROOT/'private/jobagg/output/all_jobs.sqlite3')
    p.add_argument('--profile',type=Path,default=ROOT/'private/inputs/application_context.md')
    p.add_argument('--feedback',type=Path,help='Optional approved-feedback file; only APPROVED_UPDATES is read')
    p.add_argument('--profile-overrides',type=Path,help='Optional JSON: languages, citizenships, national_countries, work_authorized_countries, disabled_families')
    p.add_argument('--as-of',help='Timezone-aware ISO time, for reproducible historical checks')
    p.add_argument('--output',type=Path,default=None)
    args=p.parse_args(argv)
    try:
        now=datetime.fromisoformat(args.as_of.replace('Z','+00:00')) if args.as_of else datetime.now(UTC)
        if now.tzinfo is None:raise ValueError('--as-of must include a timezone')
        now=now.astimezone(UTC)
        output=args.output or ROOT/'private/outputs/job_shortlists'/f'shortlist_{now:%Y%m%dT%H%M%SZ}.xlsx'
        profile=read_profile(args.profile,args.feedback,args.profile_overrides)
        rows=snapshot(args.database);records,excluded=screen(rows,profile,now)
        export_xlsx(records,profile,now,excluded,output)
        report={'as_of':now.isoformat(),'database_open_rows':len(rows),'included':len(records),
                'sheets':dict(Counter(r['group'] for r in records)),'excluded':excluded,
                'profile_evidence_sha256':profile['evidence_sha256'],'output':str(output.resolve()),
                'mode':'deterministic_python_no_network_no_llm'}
        output.with_suffix('.json').write_text(json.dumps(report,indent=2)+'\n')
        print(json.dumps(report,indent=2));return 0
    except (OSError,ValueError,sqlite3.Error,RuntimeError) as exc:
        print('Error: '+str(exc),file=sys.stderr);return 2


if __name__=='__main__':raise SystemExit(main())
