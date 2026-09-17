import importlib.util
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

PATH=Path(__file__).resolve().parents[1]/'agents/apex/scripts/shortlist_jobs.py'
spec=importlib.util.spec_from_file_location('shortlist',PATH)
app=importlib.util.module_from_spec(spec);spec.loader.exec_module(app)
NOW=datetime(2026,9,17,17,tzinfo=timezone.utc)


class ShortlistTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.profile={'families':{name:'Applicant experience: '+name for name in app.FAMILIES},
                      'languages':{'english':'full professional','japanese':'native','french':'basic'},
                      'citizenships':[],'national_countries':['JP'],'work_authorized_countries':[],'evidence_sha256':'test'}

    def tearDown(self):self.tmp.cleanup()

    def row(self,**changes):
        r={'job_key':'example:1','source_id':'example','org_id':'Example','title':'Programme Officer',
           'status':'open','closes_at':'2026-09-20T12:00:00Z','apply_url':'https://example.org/job/1',
           'description':('Support programme management and coordination of multiple projects. '*10)+' Minimum five years of programme management experience is required. Fluency in English is required.',
           'location':'Geneva, Switzerland','grade_code':'P3','source_listed_current':1,'last_seen_at':'2026-09-17T12:00:00Z'}
        r.update(changes);return r

    def test_only_candidate_and_approved_evidence(self):
        p=self.root/'profile.md';p.write_text('## USER_JOB_HISTORY_TEXT\nProgramme coordination.\nLanguages: English (full professional); French (basic).\n## JOB_DESCRIPTION_TEXT\nRequires Python data management and cash transfer experience.\n')
        f=self.root/'feedback.md';f.write_text('## APPROVED_UPDATES\nGrant oversight.\n## UPDATES_REQUIRING_CONFIRMATION\nHOPE and cash transfer work.\n')
        profile=app.read_profile(p,f)
        self.assertIn('Programme management',profile['families']);self.assertIn('Grants and partnerships',profile['families'])
        self.assertNotIn('Social protection and cash',profile['families']);self.assertNotIn('Data and information management',profile['families'])
        self.assertEqual(profile['languages']['french'],'basic')

    def test_closed_expired_stale_and_unlisted_excluded(self):
        for change in [{'status':'closed'},{'closes_at':'2026-09-16T23:59:00Z'},{'stale_current':1},{'source_listed_current':0},{'duplicate_of_job_key':'other'}]:
            with self.subTest(change=change):self.assertIsNone(app.evaluate(self.row(**change),self.profile,NOW)[0])

    def test_unknown_and_same_day_date_only_excluded(self):
        for date in [None,'invalid','2026-09-17']:
            self.assertIsNone(app.closing(self.row(closes_at=date),NOW)[0])
        parsed,error=app.closing(self.row(closes_at='2026-09-18'),NOW)
        self.assertIsNone(error);self.assertIn('Date only',parsed['precision'])

    def test_timezone_and_future_exact_today(self):
        parsed,error=app.closing(self.row(closes_at=None,closes_at_local='2026-09-17T23:59',closes_tz='Africa/Nairobi'),NOW)
        self.assertEqual(parsed['sort'].hour,20);self.assertIsNone(error)
        self.assertIsNone(app.closing(self.row(closes_at='2026-09-17T18:00',closes_tz=None),NOW)[0])

    def test_remote_national_outside_japan_is_excluded(self):
        row=self.row(title='National Monitoring and Evaluation Consultant',location='Remote',grade_code='CON')
        self.assertEqual(app.evaluate(row,self.profile,NOW)[1],'national_local_outside_japan')

    def test_national_suffix_and_embedded_national_project_officer(self):
        for title in ['Partnership Consultant, India, Remote with travel, Nationals',
                      'CFA 2026 004 | National Project Officer | Emergency Response | Roster only']:
            self.assertEqual(app.evaluate(self.row(title=title, grade_code='CON'), self.profile, NOW)[1],
                             'national_local_outside_japan')

    def test_migrant_is_not_grant_and_technical_training_excluded(self):
        self.assertEqual(app.evaluate(self.row(title='Cyber Threat Advisor – Migrant Smuggling'), self.profile, NOW)[1],
                         'no_profile_title_overlap')
        self.assertEqual(app.evaluate(self.row(title='Roster-Aircraft Maintenance Training Expert'), self.profile, NOW)[1],
                         'specialist_or_career_stage_mismatch')

    def test_subject_national_economists_not_national_recruitment(self):
        r=self.row(title='Programme Manager - Sustainable Health Investment through National Economists',location='Nairobi',grade_code='CON')
        result,reason=app.evaluate(r,self.profile,NOW)
        self.assertIsNone(reason);self.assertEqual(result['scope'],'Recruitment scope unconfirmed')

    def test_japan_national_role_included_with_citizenship_check(self):
        r=self.row(title='Administrative Officer',location='Tokyo, Japan',grade_code='NOA')
        result,_=app.evaluate(r,self.profile,NOW)
        self.assertEqual(result['group'],'Review eligibility');self.assertIn('citizenship',result['checks'])

    def test_real_remote_and_hybrid_not_confused(self):
        r,_=app.evaluate(self.row(title='Evaluation Consultant (Home-based)',grade_code='CON'),self.profile,NOW)
        self.assertEqual(r['location'],'Remote')
        r,_=app.evaluate(self.row(description=self.row()['description']+' Hybrid work and remote field locations.'),self.profile,NOW)
        self.assertNotEqual(r['location'],'Remote')

    def test_language_requirements_and_alternatives(self):
        for text,blocked in [('Fluency in English and French is required.',True),('Fluency in English or French is required.',False),('Working knowledge of French is required.',True),('Limited knowledge of French is required.',False),('French is desirable.',False)]:
            with self.subTest(text=text):
                gaps,_=app.language_checks(text,self.profile);self.assertEqual(bool(gaps),blocked)

    def test_grade_classifier_codes_not_p_grades(self):
        r,_=app.evaluate(self.row(grade_code='EXP10'),self.profile,NOW);self.assertEqual(r['grade'],'Not published')
        r,_=app.evaluate(self.row(title='Programme Management Consultant',grade_code='P4'),self.profile,NOW);self.assertEqual(r['grade'],'Consultant (ungraded)')

    def test_no_inferred_match_without_profile_evidence(self):
        self.profile['families']={}
        self.assertEqual(app.evaluate(self.row(),self.profile,NOW)[1],'no_profile_title_overlap')

    def test_sorting_deduplication_and_distinct_identical_titles(self):
        rows=[self.row(job_key='b',closes_at='2026-09-22T00:00:00Z'),self.row(job_key='a'),self.row(job_key='c'),self.row(job_key='a')]
        records,counts=app.screen(rows,self.profile,NOW)
        self.assertEqual([r['key'] for r in records],['a','c','b']);self.assertEqual(counts['duplicate'],1)

    def test_readonly_database_query_and_missing_schema(self):
        dbfile=self.root/'jobs.sqlite3'
        with sqlite3.connect(dbfile) as db:
            db.execute('CREATE TABLE jobs(job_key text,title text,status text,description text,apply_url text)')
            db.execute('INSERT INTO jobs VALUES(?,?,?,?,?)',('1','Programme Officer','open','Details','https://example.org'))
            db.execute('INSERT INTO jobs VALUES(?,?,?,?,?)',('2','Closed','closed','Details','https://example.org'))
        before=dbfile.read_bytes();rows=app.snapshot(dbfile)
        self.assertEqual([r['job_key'] for r in rows],['1']);self.assertEqual(before,dbfile.read_bytes())

    def test_excel_roundtrip_dates_links_and_no_formula_injection(self):
        from openpyxl import load_workbook
        rows=[self.row(title='=Programme Officer',description=self.row()['description']),self.row(job_key='2',title='Evaluation Consultant roster',location='Home-based',grade_code='CON')]
        records,excluded=app.screen(rows,self.profile,NOW);output=self.root/'out.xlsx'
        app.export_xlsx(records,self.profile,NOW,excluded,output)
        book=load_workbook(output)
        found=0
        for sh in list(book)[:3]:
            self.assertEqual(sh.freeze_panes,'D5')
            for row in sh.iter_rows(min_row=5):
                if isinstance(row[0].value,datetime):
                    found+=1;self.assertEqual(row[9].hyperlink.target,'https://example.org/job/1')
                    self.assertNotEqual(row[2].data_type,'f')
            self.assertFalse(any(c.data_type=='f' for row in sh for c in row))
        self.assertEqual(found,2);book.close()


if __name__=='__main__':unittest.main()
